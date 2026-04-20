"""Shared metadata helpers for supported V1 Office and PDF formats."""

from __future__ import annotations

from os import PathLike
from pathlib import Path
import shutil
import xml.etree.ElementTree as ET
import zipfile

import fitz
from docx import Document
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from pptx import Presentation

__all__ = ["clear_metadata", "read_metadata"]

_SUPPORTED_EXTENSIONS = frozenset({"xlsx", "xlsm", "docx", "pptx", "pdf"})
_SUPPORTED_EXTENSIONS_TEXT = ", ".join(sorted(_SUPPORTED_EXTENSIONS))
_OFFICE_EXTENSIONS = frozenset({"xlsx", "xlsm", "docx", "pptx"})
_OFFICE_CORE_XML_PATH = "docProps/core.xml"
_OFFICE_FIELD_TO_XML_NAME = {
    "title": "title",
    "subject": "subject",
    "creator": "creator",
    "keywords": "keywords",
    "description": "description",
    "language": "language",
    "last_modified_by": "lastModifiedBy",
    "category": "category",
    "content_status": "contentStatus",
    "identifier": "identifier",
    "version": "version",
    "revision": "revision",
    "created": "created",
    "modified": "modified",
    "last_printed": "lastPrinted",
}
_PDF_FIELD_ORDER = (
    "title",
    "author",
    "subject",
    "keywords",
    "creator",
    "producer",
    "creation_date",
    "modification_date",
    "trapped",
)
_PDF_FIELD_TO_METADATA_KEY = {
    "title": "title",
    "author": "author",
    "subject": "subject",
    "keywords": "keywords",
    "creator": "creator",
    "producer": "producer",
    "creation_date": "creationDate",
    "modification_date": "modDate",
    "trapped": "trapped",
}
_OFFICE_STRING_PROPERTY_MAP = {
    "xlsx": {
        "title": "title",
        "subject": "subject",
        "creator": "creator",
        "keywords": "keywords",
        "description": "description",
        "language": "language",
        "last_modified_by": "lastModifiedBy",
        "category": "category",
        "content_status": "contentStatus",
        "identifier": "identifier",
        "version": "version",
        "revision": "revision",
    },
    "xlsm": {
        "title": "title",
        "subject": "subject",
        "creator": "creator",
        "keywords": "keywords",
        "description": "description",
        "language": "language",
        "last_modified_by": "lastModifiedBy",
        "category": "category",
        "content_status": "contentStatus",
        "identifier": "identifier",
        "version": "version",
        "revision": "revision",
    },
    "docx": {
        "title": "title",
        "subject": "subject",
        "creator": "author",
        "keywords": "keywords",
        "description": "comments",
        "language": "language",
        "last_modified_by": "last_modified_by",
        "category": "category",
        "content_status": "content_status",
        "identifier": "identifier",
        "version": "version",
    },
    "pptx": {
        "title": "title",
        "subject": "subject",
        "creator": "author",
        "keywords": "keywords",
        "description": "comments",
        "language": "language",
        "last_modified_by": "last_modified_by",
        "category": "category",
        "content_status": "content_status",
        "identifier": "identifier",
        "version": "version",
    },
}
_OFFICE_NAMESPACE_MAP = {
    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    "dc": "http://purl.org/dc/elements/1.1/",
    "dcterms": "http://purl.org/dc/terms/",
    "dcmitype": "http://purl.org/dc/dcmitype/",
    "xsi": "http://www.w3.org/2001/XMLSchema-instance",
}

for _prefix, _uri in _OFFICE_NAMESPACE_MAP.items():
    ET.register_namespace(_prefix, _uri)


def read_metadata(file_path: str | PathLike[str] | Path) -> dict:
    """Return normalized metadata for a supported Office or PDF file."""
    source = _validate_source_path(file_path, label="Metadata source file")
    extension = _path_extension(source)
    warnings: list[str] = []

    if extension in {"xlsx", "xlsm"}:
        fields = _read_excel_metadata(source)
    elif extension == "docx":
        fields = _read_docx_metadata(source)
    elif extension == "pptx":
        fields = _read_pptx_metadata(source)
    else:
        fields = _read_pdf_metadata(source)

    return {
        "file_path": str(source),
        "extension": extension,
        "fields": fields,
        "warnings": warnings,
    }


def clear_metadata(file_path: str | PathLike[str] | Path) -> None:
    """Clear supported metadata in place for a supported Office or PDF file."""
    source = _validate_source_path(file_path, label="Metadata source file")
    extension = _path_extension(source)

    if extension in {"xlsx", "xlsm"}:
        _clear_excel_metadata(source)
        return
    if extension == "docx":
        _clear_docx_metadata(source)
        return
    if extension == "pptx":
        _clear_pptx_metadata(source)
        return
    _clear_pdf_metadata(source)


def _read_excel_metadata(path: Path) -> dict[str, str | None]:
    workbook = _load_excel_workbook(path)
    try:
        properties = workbook.properties
        fields = {
            "title": _normalize_metadata_value(properties.title),
            "subject": _normalize_metadata_value(properties.subject),
            "creator": _normalize_metadata_value(properties.creator),
            "keywords": _normalize_metadata_value(properties.keywords),
            "description": _normalize_metadata_value(properties.description),
            "language": _normalize_metadata_value(properties.language),
            "last_modified_by": _normalize_metadata_value(properties.lastModifiedBy),
            "category": _normalize_metadata_value(properties.category),
            "content_status": _normalize_metadata_value(properties.contentStatus),
            "identifier": _normalize_metadata_value(properties.identifier),
            "version": _normalize_metadata_value(properties.version),
            "revision": _normalize_metadata_value(properties.revision),
            "created": _normalize_metadata_value(properties.created),
            "modified": _normalize_metadata_value(properties.modified),
            "last_printed": _normalize_metadata_value(properties.lastPrinted),
        }
    finally:
        _close_excel_workbook(workbook)

    return _apply_office_core_xml_presence(path, fields)


def _read_docx_metadata(path: Path) -> dict[str, str | None]:
    document = Document(path)
    properties = document.core_properties
    fields = {
        "title": _normalize_metadata_value(properties.title),
        "subject": _normalize_metadata_value(properties.subject),
        "creator": _normalize_metadata_value(properties.author),
        "keywords": _normalize_metadata_value(properties.keywords),
        "description": _normalize_metadata_value(properties.comments),
        "language": _normalize_metadata_value(properties.language),
        "last_modified_by": _normalize_metadata_value(properties.last_modified_by),
        "category": _normalize_metadata_value(properties.category),
        "content_status": _normalize_metadata_value(properties.content_status),
        "identifier": _normalize_metadata_value(properties.identifier),
        "version": _normalize_metadata_value(properties.version),
        "revision": _normalize_metadata_value(properties.revision),
        "created": _normalize_metadata_value(properties.created),
        "modified": _normalize_metadata_value(properties.modified),
        "last_printed": _normalize_metadata_value(properties.last_printed),
    }
    return _apply_office_core_xml_presence(path, fields)


def _read_pptx_metadata(path: Path) -> dict[str, str | None]:
    presentation = Presentation(path)
    properties = presentation.core_properties
    fields = {
        "title": _normalize_metadata_value(properties.title),
        "subject": _normalize_metadata_value(properties.subject),
        "creator": _normalize_metadata_value(properties.author),
        "keywords": _normalize_metadata_value(properties.keywords),
        "description": _normalize_metadata_value(properties.comments),
        "language": _normalize_metadata_value(properties.language),
        "last_modified_by": _normalize_metadata_value(properties.last_modified_by),
        "category": _normalize_metadata_value(properties.category),
        "content_status": _normalize_metadata_value(properties.content_status),
        "identifier": _normalize_metadata_value(properties.identifier),
        "version": _normalize_metadata_value(properties.version),
        "revision": _normalize_metadata_value(properties.revision),
        "created": _normalize_metadata_value(properties.created),
        "modified": _normalize_metadata_value(properties.modified),
        "last_printed": _normalize_metadata_value(properties.last_printed),
    }
    return _apply_office_core_xml_presence(path, fields)


def _read_pdf_metadata(path: Path) -> dict[str, str | None]:
    document = fitz.open(path)
    try:
        _reject_encrypted_pdf(document, path)
        metadata = document.metadata or {}
        return {
            field_name: _normalize_metadata_value(metadata.get(metadata_key))
            for field_name, metadata_key in _PDF_FIELD_TO_METADATA_KEY.items()
        }
    finally:
        document.close()


def _clear_excel_metadata(path: Path) -> None:
    workbook = _load_excel_workbook(path)
    try:
        _blank_string_properties(
            workbook.properties,
            property_names=_OFFICE_STRING_PROPERTY_MAP[_path_extension(path)].values(),
        )
        workbook.save(path)
    except Exception as exc:  # pragma: no cover - library-specific failures vary by file.
        raise ValueError(f"Failed to clear Excel metadata in '{path}': {exc}") from exc
    finally:
        _close_excel_workbook(workbook)

    _strip_office_core_xml_fields(path)


def _clear_docx_metadata(path: Path) -> None:
    document = Document(path)
    try:
        _blank_string_properties(
            document.core_properties,
            property_names=_OFFICE_STRING_PROPERTY_MAP["docx"].values(),
        )
        document.save(path)
    except Exception as exc:  # pragma: no cover - library-specific failures vary by file.
        raise ValueError(f"Failed to clear Word metadata in '{path}': {exc}") from exc

    _strip_office_core_xml_fields(path)


def _clear_pptx_metadata(path: Path) -> None:
    presentation = Presentation(path)
    try:
        _blank_string_properties(
            presentation.core_properties,
            property_names=_OFFICE_STRING_PROPERTY_MAP["pptx"].values(),
        )
        presentation.save(path)
    except Exception as exc:  # pragma: no cover - library-specific failures vary by file.
        raise ValueError(f"Failed to clear PowerPoint metadata in '{path}': {exc}") from exc

    _strip_office_core_xml_fields(path)


def _clear_pdf_metadata(path: Path) -> None:
    document = fitz.open(path)
    temp_output = _temporary_output_path(path)
    try:
        _reject_encrypted_pdf(document, path)
        document.set_metadata({})
        document.save(temp_output, garbage=4, deflate=True)
    except Exception as exc:  # pragma: no cover - library-specific failures vary by file.
        if temp_output.exists():
            temp_output.unlink()
        raise ValueError(f"Failed to clear PDF metadata in '{path}': {exc}") from exc
    finally:
        if not document.is_closed:
            document.close()

    _replace_file(temp_output, path)


def _blank_string_properties(target, *, property_names) -> None:
    for property_name in property_names:
        setattr(target, property_name, "")


def _apply_office_core_xml_presence(
    path: Path,
    fields: dict[str, str | None],
) -> dict[str, str | None]:
    raw_core_values = _read_office_core_xml_values(path)
    normalized: dict[str, str | None] = {}
    for field_name, xml_name in _OFFICE_FIELD_TO_XML_NAME.items():
        if xml_name not in raw_core_values:
            normalized[field_name] = None
            continue
        raw_value = raw_core_values[xml_name]
        normalized[field_name] = fields[field_name] if raw_value not in (None, "") else None
    return normalized


def _read_office_core_xml_values(path: Path) -> dict[str, str | None]:
    try:
        with zipfile.ZipFile(path) as archive:
            try:
                raw_xml = archive.read(_OFFICE_CORE_XML_PATH)
            except KeyError:
                return {}
    except zipfile.BadZipFile as exc:
        raise ValueError(f"Office file '{path}' is not a valid Open XML package.") from exc

    try:
        root = ET.fromstring(raw_xml)
    except ET.ParseError as exc:
        raise ValueError(f"Office file '{path}' has unreadable core metadata XML.") from exc

    values: dict[str, str | None] = {}
    for child in root:
        values[_xml_local_name(child.tag)] = child.text
    return values


def _strip_office_core_xml_fields(path: Path) -> None:
    temp_output = _temporary_output_path(path)
    try:
        with zipfile.ZipFile(path) as source_archive, zipfile.ZipFile(temp_output, "w") as target_archive:
            for info in source_archive.infolist():
                payload = source_archive.read(info.filename)
                if info.filename == _OFFICE_CORE_XML_PATH:
                    payload = _strip_core_xml_payload(payload, source=path)
                target_archive.writestr(info, payload)
    except Exception:
        if temp_output.exists():
            temp_output.unlink()
        raise

    _replace_file(temp_output, path)


def _strip_core_xml_payload(payload: bytes, *, source: Path) -> bytes:
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as exc:
        raise ValueError(f"Office file '{source}' has unreadable core metadata XML.") from exc

    for child in list(root):
        if _xml_local_name(child.tag) in set(_OFFICE_FIELD_TO_XML_NAME.values()):
            root.remove(child)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _load_excel_workbook(path: Path) -> Workbook:
    try:
        return load_workbook(path, keep_vba=_path_extension(path) == "xlsm")
    except Exception as exc:  # pragma: no cover - library-specific failures vary by file.
        raise ValueError(f"Failed to load Excel workbook '{path}': {exc}") from exc


def _close_excel_workbook(workbook: Workbook) -> None:
    vba_archive = getattr(workbook, "vba_archive", None)
    if vba_archive is not None:
        vba_archive.close()
    close = getattr(workbook, "close", None)
    if callable(close):
        close()


def _reject_encrypted_pdf(document, path: Path) -> None:
    if getattr(document, "needs_pass", False):
        raise NotImplementedError(f"Encrypted PDF files are outside metadata V1 scope: '{path}'.")


def _coerce_path(value: str | PathLike[str] | Path) -> Path:
    return Path(value)


def _validate_source_path(file_path: str | PathLike[str] | Path, *, label: str) -> Path:
    path = _coerce_path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"{label} '{path}' does not exist.")
    if path.is_symlink() or not path.is_file():
        raise ValueError(f"{label} '{path}' is not a regular file.")

    extension = _path_extension(path)
    if extension not in _SUPPORTED_EXTENSIONS:
        display_extension = extension or "<none>"
        raise ValueError(
            f"{label} '{path}' has unsupported extension '{display_extension}'. "
            f"Supported extensions: {_SUPPORTED_EXTENSIONS_TEXT}."
        )
    return path


def _path_extension(path: Path) -> str:
    return path.suffix.lower().lstrip(".")


def _normalize_metadata_value(value) -> str | None:
    if value in (None, ""):
        return None
    isoformat = getattr(value, "isoformat", None)
    if callable(isoformat):
        return isoformat()
    text = str(value).strip()
    return text or None


def _xml_local_name(tag: str) -> str:
    if "}" not in tag:
        return tag
    return tag.rsplit("}", 1)[1]


def _temporary_output_path(path: Path) -> Path:
    return path.with_name(f".{path.stem}.metadata-tmp{path.suffix}")


def _replace_file(temp_output: Path, final_output: Path) -> None:
    try:
        shutil.move(str(temp_output), str(final_output))
    finally:
        if temp_output.exists():
            temp_output.unlink()
